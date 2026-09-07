# -*- coding: utf-8 -*-
"""Markdown -> Word for the Task 3 build documents.

Keeps the .md files as the single source of truth; this renders them as styled .docx.
Handles: headings, paragraphs, bullets (incl. nesting), numbered lists, pipe tables,
blockquotes, fenced code blocks, and inline **bold** / *italic* / `code` / [links](url).
"""
import re, os, sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ACCENT = RGBColor(0x0B, 0x5D, 0x66)
ACCENT_D = RGBColor(0x08, 0x4A, 0x52)
INK = RGBColor(0x13, 0x1A, 0x1C)
MUTED = RGBColor(0x5A, 0x68, 0x6C)
CRIT = RGBColor(0xA6, 0x3A, 0x32)
HDR_FILL = "EAF1F1"
ZEBRA = "F8FAFA"
BODY_FONT = "Calibri"
MONO_FONT = "Consolas"


# --------------------------- low-level helpers ---------------------------
def shade(cell, hexfill):
    tcPr = cell._tc.get_or_add_tcPr()
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:color"), "auto")
    el.set(qn("w:fill"), hexfill)
    tcPr.append(el)


def set_cell_margins(table, top=60, bottom=60, left=100, right=100):
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for tag, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement("w:" + tag)
        e.set(qn("w:w"), str(val))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tblPr.append(mar)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def _field(paragraph, code, size=7.5):
    f1 = OxmlElement("w:fldChar")
    f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = code
    f2 = OxmlElement("w:fldChar")
    f2.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    run.font.size = Pt(size)
    run.font.color.rgb = MUTED
    run.font.name = BODY_FONT
    run._r.append(f1)
    run._r.append(it)
    run._r.append(f2)


def add_footer(section, text):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text + "    ")
    r.font.size = Pt(7.5)
    r.font.color.rgb = MUTED
    r.font.name = BODY_FONT
    _field(p, "PAGE")
    r2 = p.add_run(" / ")
    r2.font.size = Pt(7.5)
    r2.font.color.rgb = MUTED
    r2.font.name = BODY_FONT
    _field(p, "NUMPAGES")


# --------------------------- inline formatting ---------------------------
INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]*\)|(?<!\*)\*(?!\*)[^*]+\*(?!\*))")


def add_inline(paragraph, text, size=10, color=None, bold_all=False, italic_all=False):
    """Append `text` to `paragraph`, honouring markdown inline marks."""
    text = text.replace("\\*", "").replace("<br>", " ")
    for part in INLINE.split(text):
        if not part:
            continue
        is_code = False
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.font.name = MONO_FONT
            run.font.color.rgb = ACCENT_D
            run.font.size = Pt(size - 0.5)
            is_code = True
        elif part.startswith("[") and "](" in part:
            run = paragraph.add_run(part[1:part.index("](")])
            run.font.color.rgb = ACCENT
            run.underline = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(part)
        if not is_code:
            run.font.name = BODY_FONT
            run.font.size = Pt(size)
            if color is not None:
                run.font.color.rgb = color
        if bold_all:
            run.bold = True
        if italic_all:
            run.italic = True


# --------------------------- document chrome ---------------------------
def new_document(title, subtitle, eyebrow):
    doc = Document()
    doc.core_properties.title = title
    doc.core_properties.subject = "Brightmoney APM KYC case study - Task 3 build documentation"
    doc.core_properties.category = "Build documentation"

    s = doc.sections[0]
    s.page_width, s.page_height = Inches(8.27), Inches(11.69)  # A4
    s.left_margin = s.right_margin = Inches(0.85)
    s.top_margin = Inches(0.8)
    s.bottom_margin = Inches(0.75)

    n = doc.styles["Normal"]
    n.font.name = BODY_FONT
    n.font.size = Pt(10)
    n.font.color.rgb = INK
    n.paragraph_format.space_after = Pt(7)
    n.paragraph_format.line_spacing = 1.15
    n._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)

    for name, size, color, before, after in (
            ("Heading 1", 15, ACCENT_D, 20, 6),
            ("Heading 2", 11.5, INK, 14, 4),
            ("Heading 3", 10.5, MUTED, 11, 3)):
        st = doc.styles[name]
        st.font.name = BODY_FONT
        st.font.size = Pt(size)
        st.font.color.rgb = color
        st.font.bold = True
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(eyebrow.upper())
    r.font.name = BODY_FONT
    r.font.size = Pt(7.5)
    r.bold = True
    r.font.color.rgb = ACCENT

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    r.font.name = BODY_FONT
    r.font.size = Pt(24)
    r.bold = True
    r.font.color.rgb = INK

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(13)
    r = p.add_run(subtitle)
    r.font.name = BODY_FONT
    r.font.size = Pt(10.5)
    r.font.color.rgb = MUTED
    r.italic = True

    add_footer(s, title)
    return doc


def hrule(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(9)
    pPr = p._p.get_or_add_pPr()
    bd = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), "6")
    b.set(qn("w:space"), "1")
    b.set(qn("w:color"), "D8E1E2")
    bd.append(b)
    pPr.append(bd)


# --------------------------- markdown parsing ---------------------------
def split_row(line):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def is_sep(line):
    return bool(re.match(r"^\s*\|?[\s:\-|]+\|[\s:\-|]*$", line)) and "-" in line


def add_table(doc, rows, aligns):
    ncols = max(len(r) for r in rows)
    t = doc.add_table(rows=0, cols=ncols)
    t.style = "Table Grid"
    t.autofit = True
    set_cell_margins(t)

    hdr = t.add_row()
    repeat_header(hdr)
    for i, txt in enumerate(rows[0] + [""] * (ncols - len(rows[0]))):
        c = hdr.cells[i]
        shade(c, HDR_FILL)
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        if aligns[i] == "r":
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        add_inline(p, txt, size=8, color=ACCENT_D, bold_all=True)

    for n, row in enumerate(rows[1:]):
        tr = t.add_row()
        for i, txt in enumerate(row + [""] * (ncols - len(row))):
            c = tr.cells[i]
            if n % 2 == 1:
                shade(c, ZEBRA)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            if aligns[i] == "r":
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            add_inline(p, txt, size=8.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def convert(md_path, docx_path, title, subtitle, eyebrow, skip_until=None):
    lines = open(md_path, encoding="utf-8").read().split("\n")
    doc = new_document(title, subtitle, eyebrow)

    i = 0
    started = skip_until is None
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        stripped = line.strip()

        if not started:
            if skip_until in stripped:
                started = True
            else:
                i += 1
                continue

        # fenced code block
        if stripped.startswith("```"):
            i += 1
            block = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1
            for bl in block:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                r = p.add_run(bl if bl.strip() else " ")
                r.font.name = MONO_FONT
                r.font.size = Pt(7)
                r.font.color.rgb = INK
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
            continue

        # table
        if stripped.startswith("|") and i + 1 < len(lines) and is_sep(lines[i + 1]):
            header = split_row(stripped)
            aligns = ["r" if c.strip().endswith(":") and c.strip().startswith("-") is False
                      or c.strip().startswith("-") and c.strip().endswith(":")
                      else "l" for c in split_row(lines[i + 1])]
            aligns += ["l"] * (len(header) - len(aligns))
            rows = [header]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i]))
                i += 1
            add_table(doc, rows, aligns)
            continue

        # blank
        if not stripped:
            i += 1
            continue

        # horizontal rule
        if re.match(r"^-{3,}$", stripped) or re.match(r"^\*{3,}$", stripped):
            hrule(doc)
            i += 1
            continue

        # headings
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level, text = len(m.group(1)), m.group(2)
            if level == 1:
                i += 1
                continue  # title already rendered in the masthead
            style = "Heading 1" if level == 2 else ("Heading 2" if level == 3 else "Heading 3")
            p = doc.add_paragraph(style=style)
            sz = 15 if style == "Heading 1" else (11.5 if style == "Heading 2" else 10.5)
            col = ACCENT_D if style == "Heading 1" else (INK if style == "Heading 2" else MUTED)
            add_inline(p, text, size=sz, color=col, bold_all=True)
            i += 1
            continue

        # blockquote
        if stripped.startswith(">"):
            buf = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip().lstrip(">").strip())
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.22)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(8)
            pPr = p._p.get_or_add_pPr()
            bd = OxmlElement("w:pBdr")
            lft = OxmlElement("w:left")
            lft.set(qn("w:val"), "single")
            lft.set(qn("w:sz"), "12")
            lft.set(qn("w:space"), "8")
            lft.set(qn("w:color"), "0B5D66")
            bd.append(lft)
            pPr.append(bd)
            add_inline(p, " ".join(buf), size=9.5, color=MUTED)
            continue

        # list item (bullet or numbered), with continuation lines
        m = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", raw)
        if m:
            indent, marker, text = m.group(1), m.group(2), m.group(3)
            i += 1
            while i < len(lines):
                nxt = lines[i]
                if (not nxt.strip() or re.match(r"^(\s*)([-*]|\d+\.)\s+", nxt)
                        or nxt.strip().startswith("#") or nxt.strip().startswith("|")
                        or re.match(r"^-{3,}$", nxt.strip())):
                    break
                text += " " + nxt.strip()
                i += 1
            numbered = marker not in ("-", "*")
            style = "List Number" if numbered else "List Bullet"
            lvl = len(indent) // 2
            p = doc.add_paragraph(style=style)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.12
            p.paragraph_format.left_indent = Inches(0.25 + 0.22 * lvl)
            add_inline(p, text, size=10)
            continue

        # paragraph, with continuation lines
        text = stripped
        i += 1
        while i < len(lines):
            nxt = lines[i]
            if (not nxt.strip() or nxt.strip().startswith("#") or nxt.strip().startswith("|")
                    or nxt.strip().startswith(">") or nxt.strip().startswith("```")
                    or re.match(r"^(\s*)([-*]|\d+\.)\s+", nxt)
                    or re.match(r"^-{3,}$", nxt.strip())):
                break
            text += " " + nxt.strip()
            i += 1
        p = doc.add_paragraph()
        add_inline(p, text, size=10)

    doc.save(docx_path)
    return docx_path
